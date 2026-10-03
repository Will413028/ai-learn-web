import io
from pathlib import Path
from zipfile import ZipFile

from docx import Document as WordDocument
from pypdf import PdfReader

from app.providers import ProviderError
from app.settings import Settings

FORMATS = {".pdf", ".docx", ".md", ".txt"}


def parse_document(filename: str, data: bytes, settings: Settings) -> list[str]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".txt", ".md"}:
            text = data.decode("utf-8-sig")
        elif suffix == ".docx":
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError("expanded document too large")
            doc = WordDocument(io.BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] +
                [" ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
        elif suffix == ".pdf":
            pdf = PdfReader(io.BytesIO(data))
            if pdf.is_encrypted or len(pdf.pages) > 200:
                raise ValueError("unsupported PDF")
            pieces = []
            length = 0
            for page in pdf.pages:
                piece = page.extract_text() or ""
                length += len(piece)
                if length > settings.max_chunks * settings.chunk_size:
                    raise ValueError("document too long")
                pieces.append(piece)
            text = "\n".join(pieces)
        else:
            raise ValueError("unsupported format")
        text = text.replace("\x00", "").strip()
        if not text:
            raise ValueError("empty document")
        step = settings.chunk_size - settings.chunk_overlap
        chunks = [text[i:i + settings.chunk_size] for i in range(0, len(text), step)]
        if len(chunks) > settings.max_chunks:
            raise ValueError("too many chunks")
        return chunks
    except Exception:
        raise ProviderError("文件無法解析，請檢查格式、內容與大小後重試") from None
