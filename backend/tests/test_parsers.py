from dataclasses import replace
import io

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

from app.files import parse_document
from app.providers import ProviderError
from app.settings import Settings


def pdf_fixture() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
        DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (Python lists and dictionaries) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_fixture() -> bytes:
    doc = Document()
    doc.add_paragraph("Python lists")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "and dictionaries"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


@pytest.mark.parametrize("name,data", [("lesson.pdf", pdf_fixture()), ("lesson.docx", docx_fixture()),
    ("lesson.md", b"# Python lists and dictionaries"), ("lesson.txt", b"Python lists and dictionaries")])
def test_supported_formats_extract_actual_text(name: str, data: bytes) -> None:
    chunks = parse_document(name, data, Settings())
    text = " ".join(chunks)
    assert "Python" in text and "lists" in text and "dictionaries" in text


def test_empty_scanned_pdf_requires_text() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ProviderError):
        parse_document("scan.pdf", output.getvalue(), Settings())


def test_chunk_limit_stops_unbounded_documents() -> None:
    with pytest.raises(ProviderError):
        parse_document("large.txt", b"Python " * 30, replace(Settings(), chunk_size=20, chunk_overlap=5, max_chunks=2))


def test_chunks_overlap_without_losing_text() -> None:
    text = "abcdefghijklmnopqrstuvwx"
    chunks = parse_document("x.txt", text.encode(), replace(Settings(), chunk_size=10, chunk_overlap=2))
    assert chunks == ["abcdefghij", "ijklmnopqr", "qrstuvwx"]
